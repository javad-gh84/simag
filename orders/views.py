from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.contrib.admin.views.decorators import staff_member_required
from django.contrib import messages
from django.db.models import Sum, Count, Q
from django.utils import timezone
from datetime import timedelta
from django.http import HttpResponse
from .models import Order, OrderItem
from products.models import Product, Category
from accounts.models import User
from cart.models import Cart
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill


@login_required
def order_list(request):
    orders = Order.objects.filter(user=request.user).order_by('-created_at')
    return render(request, 'orders/order_list.html', {'orders': orders})


@login_required
def order_detail(request, order_id):
    order = get_object_or_404(Order, id=order_id, user=request.user)
    return render(request, 'orders/order_detail.html', {'order': order})


@login_required
def create_order(request):
    cart = Cart.objects.get(user=request.user)
    cart_items = cart.items.all()

    if not cart_items:
        messages.error(request, '❌ سبد خرید شما خالی است!')
        return redirect('cart:cart_detail')

    if request.method == 'POST':
        shipping_address = request.POST.get('shipping_address')
        shipping_method = request.POST.get('shipping_method', 'پست')
        shipping_cost = 25000

        if not shipping_address:
            messages.error(request, '⚠️ لطفاً آدرس ارسال را وارد کنید.')
            return redirect('orders:checkout')

        total_amount = cart.total_price
        final_amount = total_amount + shipping_cost

        from django.db import transaction
        with transaction.atomic():
            order = Order.objects.create(
                user=request.user,
                total_amount=total_amount,
                discount_amount=0,
                final_amount=final_amount,
                shipping_cost=shipping_cost,
                shipping_method=shipping_method,
                shipping_address=shipping_address,
                status='pending',
            )

            for item in cart_items:
                OrderItem.objects.create(
                    order=order,
                    product=item.product,
                    quantity=item.quantity,
                    price=item.product.price,
                    total_price=item.quantity * item.product.price
                )

            cart_items.delete()

        messages.success(request, f'✅ سفارش شما با موفقیت ثبت شد! شماره سفارش: {order.order_number}')
        return redirect('orders:order_detail', order_id=order.id)

    return redirect('cart:cart_detail')


@login_required
def checkout(request):
    cart = Cart.objects.get(user=request.user)
    cart_items = cart.items.all()

    if not cart_items:
        messages.error(request, '❌ سبد خرید شما خالی است!')
        return redirect('cart:cart_detail')

    total_price = cart.total_price
    shipping_cost = 25000
    final_price = total_price + shipping_cost

    context = {
        'cart_items': cart_items,
        'total_price': total_price,
        'shipping_cost': shipping_cost,
        'final_price': final_price,
        'user': request.user,
    }
    return render(request, 'orders/checkout.html', context)


@login_required
def cancel_order(request, order_id):
    order = get_object_or_404(Order, id=order_id, user=request.user)

    if order.status == 'pending':
        order.status = 'cancelled'
        order.save()
        messages.success(request, '✅ سفارش با موفقیت لغو شد!')
    else:
        messages.error(request, '⚠️ فقط سفارشات در انتظار قابل لغو هستند!')

    return redirect('orders:order_detail', order_id=order.id)


@login_required
def track_order(request, order_id):
    order = get_object_or_404(Order, id=order_id, user=request.user)
    return render(request, 'orders/track_order.html', {'order': order})


# ============================================================
#                      گزارشات (ادمین)
# ============================================================

@staff_member_required
def reports_dashboard(request):
    """داشبورد گزارشات"""
    
    # ===== آمار کلی =====
    total_orders = Order.objects.count()
    total_users = User.objects.count()
    total_products = Product.objects.filter(is_active=True).count()
    total_revenue = Order.objects.filter(status='delivered').aggregate(Sum('total_amount'))['total_amount__sum'] or 0
    
    # ===== فروش امروز =====
    today = timezone.now().date()
    today_orders = Order.objects.filter(created_at__date=today).count()
    today_revenue = Order.objects.filter(created_at__date=today, status='delivered').aggregate(Sum('total_amount'))['total_amount__sum'] or 0
    
    # ===== فروش هفته =====
    week_ago = today - timedelta(days=7)
    week_orders = Order.objects.filter(created_at__date__gte=week_ago).count()
    week_revenue = Order.objects.filter(created_at__date__gte=week_ago, status='delivered').aggregate(Sum('total_amount'))['total_amount__sum'] or 0
    
    # ===== فروش ماه =====
    month_ago = today - timedelta(days=30)
    month_orders = Order.objects.filter(created_at__date__gte=month_ago).count()
    month_revenue = Order.objects.filter(created_at__date__gte=month_ago, status='delivered').aggregate(Sum('total_amount'))['total_amount__sum'] or 0
    
    # ===== پرفروش‌ترین محصولات =====
    top_products = Product.objects.filter(
        is_active=True
    ).annotate(
        total_sold=Sum('orderitem__quantity')
    ).order_by('-total_sold')[:10]
    
    # ===== وضعیت سفارشات =====
    status_counts = {
        'pending': Order.objects.filter(status='pending').count(),
        'processing': Order.objects.filter(status='processing').count(),
        'shipped': Order.objects.filter(status='shipped').count(),
        'delivered': Order.objects.filter(status='delivered').count(),
        'cancelled': Order.objects.filter(status='cancelled').count(),
    }
    
    # ===== فروش ماهانه (۱۲ ماه اخیر) =====
    monthly_sales = []
    for i in range(11, -1, -1):
        date = timezone.now().date() - timedelta(days=30*i)
        month_start = date.replace(day=1)
        if i == 0:
            month_end = timezone.now().date()
        else:
            month_end = (month_start + timedelta(days=32)).replace(day=1) - timedelta(days=1)
        
        revenue = Order.objects.filter(
            status='delivered',
            created_at__date__gte=month_start,
            created_at__date__lte=month_end
        ).aggregate(Sum('total_amount'))['total_amount__sum'] or 0
        
        monthly_sales.append({
            'month': month_start.strftime('%b'),
            'revenue': revenue,
            'orders': Order.objects.filter(
                created_at__date__gte=month_start,
                created_at__date__lte=month_end
            ).count(),
        })
    
    context = {
        'total_orders': total_orders,
        'total_users': total_users,
        'total_products': total_products,
        'total_revenue': total_revenue,
        'today_orders': today_orders,
        'today_revenue': today_revenue,
        'week_orders': week_orders,
        'week_revenue': week_revenue,
        'month_orders': month_orders,
        'month_revenue': month_revenue,
        'top_products': top_products,
        'status_counts': status_counts,
        'monthly_sales': monthly_sales,
    }
    
    return render(request, 'orders/reports_dashboard.html', context)


@staff_member_required
def sales_report(request):
    """گزارش فروش با فیلتر"""
    
    date_from = request.GET.get('date_from')
    date_to = request.GET.get('date_to')
    status = request.GET.get('status')
    
    orders = Order.objects.all().order_by('-created_at')
    
    if date_from:
        orders = orders.filter(created_at__date__gte=date_from)
    if date_to:
        orders = orders.filter(created_at__date__lte=date_to)
    if status:
        orders = orders.filter(status=status)
    
    total_orders = orders.count()
    total_revenue = orders.filter(status='delivered').aggregate(Sum('total_amount'))['total_amount__sum'] or 0
    total_items = OrderItem.objects.filter(order__in=orders).aggregate(Sum('quantity'))['quantity__sum'] or 0
    
    context = {
        'orders': orders,
        'total_orders': total_orders,
        'total_revenue': total_revenue,
        'total_items': total_items,
        'date_from': date_from,
        'date_to': date_to,
        'status': status,
    }
    
    return render(request, 'orders/sales_report.html', context)


@staff_member_required
def product_report(request):
    """گزارش محصولات"""
    
    products = Product.objects.filter(is_active=True).annotate(
        total_sold=Sum('orderitem__quantity'),
        total_revenue=Sum('orderitem__total_price'),
        order_count=Count('orderitem'),
    ).order_by('-total_sold')
    
    context = {
        'products': products,
    }
    
    return render(request, 'orders/product_report.html', context)


@staff_member_required
def user_report(request):
    """گزارش کاربران"""
    
    users = User.objects.annotate(
        order_count=Count('orders'),
        total_spent=Sum('orders__total_amount', filter=Q(orders__status='delivered')),
    ).order_by('-order_count')
    
    context = {
        'users': users,
    }
    
    return render(request, 'orders/user_report.html', context)


@staff_member_required
def export_excel(request):
    """خروجی اکسل از گزارشات"""
    
    orders = Order.objects.all().order_by('-created_at')
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "گزارش سفارشات"
    
    headers = ['شماره سفارش', 'کاربر', 'مبلغ کل', 'وضعیت', 'تاریخ']
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4CAF50", end_color="4CAF50", fill_type="solid")
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
    
    for row, order in enumerate(orders, 2):
        ws.cell(row=row, column=1, value=order.order_number)
        ws.cell(row=row, column=2, value=order.user.phone)
        ws.cell(row=row, column=3, value=order.total_amount)
        ws.cell(row=row, column=4, value=order.get_status_display())
        ws.cell(row=row, column=5, value=order.created_at.strftime('%Y/%m/%d %H:%M'))
    
    for col in range(1, 6):
        ws.column_dimensions[chr(64 + col)].width = 20
    
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=orders_report.xlsx'
    wb.save(response)
    
    return response